from __future__ import annotations
import argparse, json, re
from pathlib import Path

FIG_RE = re.compile(
    r"\b(?:Figure|Fig\.?)\s+(S?\d+)"
    r"(?:\s*([A-Za-z](?:\s*[–\-]\s*[A-Za-z])?))?",
    flags=re.IGNORECASE,
)
TABLE_RE = re.compile(
    r"\b(?:Table)\s+(S?\d+)",
    flags=re.IGNORECASE,
)

def read_text(path: Path) -> str:
    suf=path.suffix.lower()
    if suf in {".txt",".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suf==".docx":
        from docx import Document
        doc=Document(path)
        chunks=[p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                chunks.extend(c.text for c in row.cells)
        return "\n".join(chunks)
    raise ValueError("Supported input: .txt, .md, .docx")

def expand_panel(token):
    if not token:
        return []
    token=re.sub(r"\s+","",token.upper())
    if "–" in token or "-" in token:
        parts=re.split(r"[–-]",token)
        if len(parts)==2 and len(parts[0])==1 and len(parts[1])==1:
            a,b=map(ord,parts)
            if a<=b:
                return [chr(x) for x in range(a,b+1)]
    return [token]

def extract(path: Path):
    text=read_text(path)
    figures=[]
    for m in FIG_RE.finditer(text):
        fig=m.group(1).upper()
        panel_token=m.group(2)
        figures.append({
            "figure":fig,
            "panel_token":panel_token.upper() if panel_token else None,
            "panels":expand_panel(panel_token),
            "matched_text":m.group(0),
            "start":m.start(),
            "end":m.end(),
        })
    tables=[{
        "table":m.group(1).upper(),
        "matched_text":m.group(0),
        "start":m.start(),
        "end":m.end(),
    } for m in TABLE_RE.finditer(text)]
    return {"file":str(path),"figures":figures,"tables":tables}

if __name__=="__main__":
    ap=argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--out")
    args=ap.parse_args()
    res=extract(Path(args.file))
    txt=json.dumps(res,indent=2,ensure_ascii=False)
    print(txt)
    if args.out:
        Path(args.out).write_text(txt,encoding="utf-8")
